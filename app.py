from flask import Flask, render_template, request, redirect, url_for, flash
import os
from werkzeug.utils import secure_filename
import PyPDF2
from docx import Document
# Replace these placeholders with your actual values
TENANT = "platformzispire"
POLICY = "B2C_1_susi"



app = Flask(__name__)
app.config['UPLOAD_FOLDER'] = 'uploads'
app.config['ALLOWED_EXTENSIONS'] = {'pdf', 'docx', 'txt'}
app.secret_key = "supersecretkey"

if not os.path.exists(app.config['UPLOAD_FOLDER']):
    os.makedirs(app.config['UPLOAD_FOLDER'])
    
def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in app.config['ALLOWED_EXTENSIONS']

@app.route('/')
def main():
    return render_template('main.html')

@app.route('/upload', methods=['GET', 'POST'])
def upload():
    if request.method == 'POST':
        files = request.files.getlist('file')
        for file in files:
            if file and allowed_file(file.filename):
                filename = secure_filename(file.filename)
                file.save(os.path.join(app.config['UPLOAD_FOLDER'], filename))
        flash('Files uploaded successfully!')
        return redirect(url_for('draft'))
    return render_template('upload.html')

@app.route('/draft')
def draft():
    files = os.listdir(app.config['UPLOAD_FOLDER'])
    sentences = []
    for file in files:
        filepath = os.path.join(app.config['UPLOAD_FOLDER'], file)
        if file.endswith('.pdf'):
            with open(filepath, 'rb') as f:
                reader = PyPDF2.PdfReader(f)
                page = reader.pages[0]
                sentences.append(page.extract_text().split('.')[0])
        elif file.endswith('.docx'):
            doc = Document(filepath)
            sentences.append(doc.paragraphs[0].text.split('.')[0])
        elif file.endswith('.txt'):
            with open(filepath, 'r') as f:
                sentences.append(f.readline().split('.')[0])
        
        # Delete the file after reading the first sentence
        os.remove(filepath)
        
    return render_template('draft.html', sentences=sentences)

@app.route('/logout')
def logout():
    redirect_uri = "https%3A%2F%2Fwww.zispire.com%2F"
    logout_url = f"https://{TENANT}.b2clogin.com/{TENANT}.onmicrosoft.com/{POLICY}/oauth2/v2.0/logout?post_logout_redirect_uri={redirect_uri}"
    return redirect(logout_url)

if __name__ == '__main__':
    app.run(debug=True,port=5000)
